import openpyxl, re
import pandas as pd
import numpy as np

CT_FILE = '/mnt/user-data/uploads/MIRCT_FILES_5683369_2020-05-05_Eclipse_Doses__i1Prst__AR__324_Dose_Table.xlsx'
MRI_FILE = '/mnt/user-data/uploads/Input_Prost_scores_latetoxicity_anyonmized.xlsx'

RX_DOSE = 40.0  # Gy, prescription (confirmed via 100% norm = 40.00 Gy in CT header)

ORGAN_MAP_CT = {
    'O_Bldr': 'Bladder',
    'O_Rctm': 'Rectum',
    'O_Femr_Lt': 'Femoral Head (L)',
    'O_Femr_Rt': 'Femoral Head (R)',
    'O_Pblb': 'Penile Bulb',
    'O_sigmd': 'Sigmoid',
    'O_Smallbwl': 'Small Bowel',
    'O_AnalCanal': 'Anal Canal',
}

def compute_metrics(doses, fracs):
    """doses: array of Gy values ascending; fracs: array of remaining-volume fraction (0-1) at each dose.
    Returns dict of Dmean, Dmax, V50pctRx (% vol >= 0.5*Rx), V90pctRx (% vol >= 0.9*Rx)."""
    doses = np.array(doses, dtype=float)
    fracs = np.array(fracs, dtype=float)
    valid = ~np.isnan(fracs)
    doses, fracs = doses[valid], fracs[valid]
    if len(doses) < 2:
        return None
    # Dmean via trapezoid integration of the cumulative DVH (volume fraction vs dose)
    dmean = np.trapezoid(fracs, doses) if hasattr(np, 'trapezoid') else np.trapz(fracs, doses)
    # Dmax: highest dose bin with fraction still > ~0.5% (avoid noise at exact 0)
    nonzero = doses[fracs > 0.005]
    dmax = nonzero.max() if len(nonzero) else 0.0
    def vol_at_dose(d):
        return float(np.interp(d, doses, fracs, left=fracs[0], right=0.0)) * 100
    return {
        'Dmean_Gy': dmean,
        'Dmax_Gy': dmax,
        'V50pctRx': vol_at_dose(0.5 * RX_DOSE),
        'V90pctRx': vol_at_dose(0.9 * RX_DOSE),
    }

# ---------- CT cohort (File 1) ----------
def parse_ct():
    wb = openpyxl.load_workbook(CT_FILE, data_only=True, read_only=True)
    records = []
    for sn in wb.sheetnames:
        ws = wb[sn]
        rows = list(ws.iter_rows(values_only=True))
        header = rows[1]  # row index 1 = 'Contour', 'Absolute Volume/ml', dose cols..., CI, nCI, HI
        # extract dose (Gy) for each dose column
        dose_cols = []  # (col_index, dose_gy)
        for i, h in enumerate(header):
            if isinstance(h, str):
                m = re.search(r'([\d.]+)\s*Gy\s*$', h)
                if m:
                    dose_cols.append((i, float(m.group(1))))
        doses = [d for _, d in dose_cols]
        for row in rows[2:]:
            contour = row[0]
            if contour not in ORGAN_MAP_CT:
                continue
            organ = ORGAN_MAP_CT[contour]
            abs_vol = row[1]
            fracs = []
            for i, _ in dose_cols:
                v = row[i]
                fracs.append(np.nan if (v is None or v == '-' or v == '∞' or isinstance(v, str)) else float(v))
            m = compute_metrics(doses, fracs)
            if m is None:
                continue
            m.update({'patient_id': sn, 'cohort': 'CT', 'organ': organ, 'abs_volume_ml': abs_vol})
            records.append(m)
    return pd.DataFrame(records)

# ---------- MRI cohort (File 2) ----------
ORGAN_MAP_MRI = {
    'O_Bldr_SIM': 'Bladder',
    'O_Rctm_SIM': 'Rectum',
    'O_Trigone_SIM': 'Trigone',
    'O_Urethra_SIM': 'Urethra',
}

def parse_mri():
    wb = openpyxl.load_workbook(MRI_FILE, data_only=True, read_only=True)
    ws = wb['DVH points']
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    patient_cols = [(i, name) for i, name in enumerate(header) if name and str(name).startswith('Pat')]

    # organize: organ -> {patient_col_index: {dose: pct}}
    data = {organ: {i: {} for i, _ in patient_cols} for organ in ORGAN_MAP_MRI}
    for row in rows[1:]:
        label = row[0]
        if not label:
            continue
        m = re.match(r'([\d.]+)Gy_(.+)', label)
        if not m:
            continue
        dose = float(m.group(1))
        organ_raw = m.group(2)
        if organ_raw not in ORGAN_MAP_MRI:
            continue
        for i, _ in patient_cols:
            v = row[i]
            if v is not None:
                data[organ_raw][i][dose] = v

    records = []
    for organ_raw, organ_name in ORGAN_MAP_MRI.items():
        for i, pname in patient_cols:
            dmap = data[organ_raw][i]
            if not dmap:
                continue
            doses_sorted = sorted(dmap.keys())
            fracs = [dmap[d] / 100.0 for d in doses_sorted]
            m = compute_metrics(doses_sorted, fracs)
            if m is None:
                continue
            m.update({'patient_id': pname, 'cohort': 'MRI', 'organ': organ_name, 'abs_volume_ml': np.nan})
            records.append(m)
    return pd.DataFrame(records)

ct_df = parse_ct()
mri_df = parse_mri()
combined = pd.concat([ct_df, mri_df], ignore_index=True)
combined.to_csv('dvh_metrics_combined.csv', index=False)
print(combined.groupby(['cohort','organ']).size())
print(combined.head())
